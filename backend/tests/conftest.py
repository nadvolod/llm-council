"""Shared pytest fixtures for backend tests."""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter


@pytest.fixture
def pdf_bytes() -> bytes:
    """A minimal valid (but text-less) PDF — pypdf can read it without crashing."""
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def docx_bytes() -> bytes:
    """A DOCX with two paragraphs."""
    doc = Document()
    doc.add_paragraph("Hello from a docx file.")
    doc.add_paragraph("Second paragraph with content.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def text_bytes() -> bytes:
    return b"Plain text content for testing."
